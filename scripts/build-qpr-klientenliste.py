#!/usr/bin/env python3
"""Erzeugt die Excel-Vorlage "QPR-Klientenliste (Teil 1a)" fuer die ambulante
Qualitaetspruefung ab 01.07.2026.

Ausfuehren:  python3 scripts/build-qpr-klientenliste.py [ZIELDATEI] [ZEILEN]
Benoetigt:   openpyxl
"""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

DATA_START = 3          # erste Eingabezeile
DATA_ROWS = 400         # vorbereitete Zeilen (ueber CLI-Argument aenderbar)
DATA_END = DATA_START + DATA_ROWS - 1

# --- Auswahllisten ----------------------------------------------------------
LISTS = {
    "Mobilitaet": ("Mobilität", [
        "uneingeschränkt", "eingeschränkt", "komplett eingeschränkt"]),
    "Kognition": ("Kognition / Kommunikation", [
        "uneingeschränkt", "eingeschränkt", "komplett eingeschränkt"]),
    "Pflegegrad": ("Pflegegrad", [
        "PG 1", "PG 2", "PG 3", "PG 4", "PG 5", "Antrag gestellt", "ohne PG"]),
    "JaNein": ("Ja / Nein", ["Ja", "Nein"]),
    "Palliativ": ("Palliativversorgung", ["nein", "AAPV", "SAPV"]),
    "Kostentraeger": ("Weitere Kostenträger", [
        "keine", "SGB XII (Sozialamt)", "Beihilfe", "Selbstzahler",
        "Unfallkasse / BG", "sonstige"]),
    "Betreuungsbereiche": ("Aufgabenbereiche Betreuung (Referenz)", [
        "Gesundheitssorge", "Aufenthaltsbestimmung", "Vermögenssorge",
        "Behördenangelegenheiten", "Wohnungsangelegenheiten",
        "Post- und Fernmeldeverkehr"]),
}

# --- Spaltendefinition ------------------------------------------------------
# (Ueberschrift, Breite, Validierungsliste|None, Format|None, Kommentar|None)
SECTIONS = [
    ("1 · Stammdaten des Klienten", "1F4E79", "D6E4F0", [
        ("Nachname", 18, None, None, None),
        ("Vorname", 16, None, None, None),
        ("Geburtsdatum", 14, None, "date", None),
        ("Straße & Hausnr.", 26, None, None, None),
        ("PLZ", 8, None, "text", None),
        ("Ort", 18, None, None, None),
        ("Telefon Klient", 18, None, "text", None),
    ]),
    ("2 · Rechtliche Betreuung & Notfallkontakte", "2E6E5B", "D7EAE3", [
        ("Rechtliche Betreuung?", 15, "JaNein", None, None),
        ("Aufgabenbereiche der Betreuung", 30, None, None,
         "Mehrfachnennung möglich, z. B.: Gesundheitssorge, "
         "Aufenthaltsbestimmung, Vermögenssorge. Siehe Blatt 'Auswahllisten'."),
        ("Betreuer:in – Name", 22, None, None, None),
        ("Betreuer:in – Anschrift", 26, None, None, None),
        ("Betreuer:in – Telefon", 18, None, "text", None),
        ("Betreuer:in – E-Mail", 24, None, None, None),
        ("Bevollmächtigte / Angehörige – Name", 24, None, None, None),
        ("Angehörige – Telefon", 18, None, "text", None),
        ("Notfallkontakt (Name & Telefon)", 26, None, None,
         "Wer wird im Notfall zuerst angerufen? Name und Nummer eintragen."),
    ]),
    ("3 · Einstufung & leistungsrechtliche Zuordnung", "7A5C1E", "F0E6CE", [
        ("Pflegegrad", 15, "Pflegegrad", None, None),
        ("SGB XI – Pflegesachleistung", 13, "JaNein", None, None),
        ("SGB XI – Kombinationsleistung", 13, "JaNein", None, None),
        ("SGB XI – Entlastungsbetrag § 45b", 13, "JaNein", None, None),
        ("SGB XI – Verhinderungspflege § 39", 13, "JaNein", None, None),
        ("SGB V – HKP § 37", 13, "JaNein", None,
         "Häusliche Krankenpflege nach § 37 SGB V (ärztliche Verordnung)."),
        ("SGB V – Leistungsarten", 30, None, None,
         "z. B. Behandlungspflege, Medikamentengabe, Wundversorgung, "
         "Injektionen, Kompressionstherapie."),
        ("Weitere Kostenträger", 20, "Kostentraeger", None, None),
    ]),
    ("4 · Diagnosen & besondere Versorgungssituationen", "8C3B2E", "F4DDD8", [
        ("Relevante Pflege- / Hauptdiagnosen", 34, None, None,
         "Für die MD-Stichprobenziehung relevante Diagnosen, keine "
         "vollständige Krankengeschichte."),
        ("Außerklinische Intensivpflege (IPReG)", 15, "JaNein", None,
         "Außerklinische Intensivpflege nach § 37c SGB V / GKV-IPReG (SKI)."),
        ("Psychiatrische HKP (pHKP)", 15, "JaNein", None, None),
        ("Palliativversorgung", 16, "Palliativ", None,
         "AAPV = allgemeine, SAPV = spezialisierte ambulante "
         "Palliativversorgung."),
    ]),
    ("5 · Funktionseinschränkungen", "5B3A78", "E4DAEF", [
        ("Mobilität", 22, "Mobilitaet", None,
         "uneingeschränkt = eigenständig, ohne Hilfsmittel\n"
         "eingeschränkt = Hilfsmittel oder Unterstützung erforderlich\n"
         "komplett eingeschränkt = bettlägerig, vollständige Immobilität"),
        ("Kognition / Kommunikation", 24, "Kognition", None,
         "uneingeschränkt = orientiert, entscheidungsfähig\n"
         "eingeschränkt = z. B. leichte bis mittelgradige Demenz\n"
         "komplett eingeschränkt = z. B. schwere Demenz, Wachkoma"),
    ]),
    ("6 · Gesamtbewertung (automatisch)", "3F3F46", "DCDCE0", [
        ("Kategorie", 11, None, "formula", None),
        ("Herleitung", 44, None, "formula", None),
    ]),
    ("Organisatorisches", "555555", "E8E8E8", [
        ("Bemerkungen", 30, None, None, None),
        ("Stand (letzte Aktualisierung)", 17, None, "date", None),
    ]),
]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Klientenliste"
    lists_ws = wb.create_sheet("Auswahllisten")
    info_ws = wb.create_sheet("Hinweise")

    # ---- Auswahllisten + benannte Bereiche --------------------------------
    for idx, (key, (label, values)) in enumerate(LISTS.items(), start=1):
        col = get_column_letter(idx)
        cell = lists_ws.cell(row=1, column=idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3F3F46")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, v in enumerate(values, start=2):
            lists_ws.cell(row=r, column=idx, value=v).border = BORDER
        lists_ws.column_dimensions[col].width = max(len(label) + 2, 24)
        wb.defined_names.add(DefinedName(
            key,
            attr_text=f"Auswahllisten!${col}$2:${col}${len(values) + 1}"))
    lists_ws.row_dimensions[1].height = 32
    lists_ws.freeze_panes = "A2"

    # ---- Kopfzeilen der Klientenliste -------------------------------------
    flat = []          # (header, width, listkey, fmt, comment, section_index)
    col = 1
    for s_idx, (title, dark, light, cols) in enumerate(SECTIONS):
        start = col
        for c in cols:
            flat.append(c + (s_idx,))
            col += 1
        end = col - 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1,
                       end_column=end)
        gh = ws.cell(row=1, column=start, value=title)
        gh.font = Font(bold=True, color="FFFFFF", size=11)
        gh.fill = PatternFill("solid", fgColor=dark)
        gh.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(start, end + 1):
            ws.cell(row=1, column=c).border = BORDER

    for i, (header, width, listkey, fmt, comment, s_idx) in enumerate(
            flat, start=1):
        _, dark, light, _ = SECTIONS[s_idx]
        cell = ws.cell(row=2, column=i, value=header)
        cell.font = Font(bold=True, size=10, color="1A1A1A")
        cell.fill = PatternFill("solid", fgColor=light)
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
        if comment:
            cell.comment = Comment(comment, "QPR-Vorlage", height=140,
                                   width=320)

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 58

    # Spaltenpositionen fuer die Formeln
    names = [f[0] for f in flat]
    col_mob = get_column_letter(names.index("Mobilität") + 1)
    col_kog = get_column_letter(names.index("Kognition / Kommunikation") + 1)
    col_kat = get_column_letter(names.index("Kategorie") + 1)
    col_her = get_column_letter(names.index("Herleitung") + 1)

    # ---- Datenzeilen ------------------------------------------------------
    grey = PatternFill("solid", fgColor="F2F2F2")
    for row in range(DATA_START, DATA_END + 1):
        for i, (header, width, listkey, fmt, comment, s_idx) in enumerate(
                flat, start=1):
            cell = ws.cell(row=row, column=i)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if fmt == "date":
                cell.number_format = "DD.MM.YYYY"
            elif fmt == "text":
                cell.number_format = "@"

        m, k = f"${col_mob}{row}", f"${col_kog}{row}"
        kat = ws.cell(row=row, column=names.index("Kategorie") + 1)
        kat.value = (
            f'=IF(OR({m}="",{k}=""),"",'
            f'IF(AND({m}="komplett eingeschränkt",'
            f'{k}="komplett eingeschränkt"),"D",'
            f'IF(OR({m}="komplett eingeschränkt",'
            f'{k}="komplett eingeschränkt"),"C",'
            f'IF(OR({m}="eingeschränkt",{k}="eingeschränkt"),"B","A"))))')
        kat.font = Font(bold=True, size=12)
        kat.alignment = Alignment(horizontal="center", vertical="center")

        ref = f"${col_kat}{row}"
        her = ws.cell(row=row, column=names.index("Herleitung") + 1)
        her.value = (
            f'=IF({ref}="","",'
            f'IF({ref}="D","Mobilität UND Kognition komplett eingeschränkt '
            f'– höchste Komplexitätsstufe",'
            f'IF({ref}="C","Mobilität oder Kognition komplett eingeschränkt",'
            f'IF({ref}="B","Teileinschränkung bei Mobilität oder Kognition",'
            f'"Mobilität und Kognition uneingeschränkt"))))')
        her.fill = grey
        her.font = Font(italic=True, size=9, color="595959")
        her.alignment = Alignment(vertical="center", wrap_text=True)

    # ---- Datenvalidierung --------------------------------------------------
    for i, (header, width, listkey, fmt, comment, s_idx) in enumerate(
            flat, start=1):
        if not listkey:
            continue
        col = get_column_letter(i)
        dv = DataValidation(type="list", formula1=f"={listkey}",
                            allow_blank=True, showDropDown=False)
        dv.error = "Bitte einen Wert aus der Liste auswählen."
        dv.errorTitle = "Ungültige Eingabe"
        dv.prompt = "Wert aus der Auswahlliste wählen"
        dv.promptTitle = header
        ws.add_data_validation(dv)
        dv.add(f"{col}{DATA_START}:{col}{DATA_END}")

    # ---- Bedingte Formatierung Kategorie ----------------------------------
    kat_range = f"{col_kat}{DATA_START}:{col_kat}{DATA_END}"
    for letter, bg, fg in [("A", "C6EFCE", "1E6B33"),
                           ("B", "FFEB9C", "8A6100"),
                           ("C", "FFCC99", "9C4A16"),
                           ("D", "FFC7CE", "9C0006")]:
        ws.conditional_formatting.add(kat_range, CellIsRule(
            operator="equal", formula=[f'"{letter}"'],
            fill=PatternFill("solid", fgColor=bg),
            font=Font(bold=True, color=fg)))

    # ---- Beispielzeile (deutlich markiert) --------------------------------
    example = {
        "Nachname": "BEISPIEL – Zeile vor Nutzung löschen",
        "Vorname": "Erika",
        "Straße & Hausnr.": "Musterweg 12",
        "PLZ": "12345",
        "Ort": "Musterstadt",
        "Telefon Klient": "0123 456789",
        "Rechtliche Betreuung?": "Ja",
        "Aufgabenbereiche der Betreuung": "Gesundheitssorge, Vermögenssorge",
        "Betreuer:in – Name": "M. Muster (Betreuungsverein)",
        "Pflegegrad": "PG 4",
        "SGB XI – Pflegesachleistung": "Ja",
        "SGB V – HKP § 37": "Ja",
        "SGB V – Leistungsarten": "Wundversorgung, Medikamentengabe",
        "Weitere Kostenträger": "keine",
        "Relevante Pflege- / Hauptdiagnosen":
            "Diabetes mellitus Typ 2, Ulcus cruris, mittelgradige Demenz",
        "Mobilität": "eingeschränkt",
        "Kognition / Kommunikation": "eingeschränkt",
    }
    warn = PatternFill("solid", fgColor="FFF2CC")
    for i, (header, *_rest) in enumerate(flat, start=1):
        cell = ws.cell(row=DATA_START, column=i)
        if header in ("Kategorie", "Herleitung"):
            continue
        if header in example:
            cell.value = example[header]
        cell.fill = warn
        cell.font = Font(italic=True, color="7F6000")
    ws.cell(row=DATA_START, column=names.index("Geburtsdatum") + 1,
            value="14.03.1948")

    # ---- Ansicht & Druck ---------------------------------------------------
    ws.freeze_panes = f"C{DATA_START}"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(flat))}{DATA_END}"
    ws.print_title_rows = "1:2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    build_info_sheet(info_ws, col_kat)

    wb.save(path)
    return path


def build_info_sheet(ws, col_kat):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for c in "CDE":
        ws.column_dimensions[c].width = 26
    ws.column_dimensions["F"].width = 40
    ws.sheet_view.showGridLines = False

    r = 2
    t = ws.cell(row=r, column=2, value="QPR-Klientenliste (Teil 1a) – "
                                      "ambulante Qualitätsprüfung ab 01.07.2026")
    t.font = Font(bold=True, size=15, color="1F4E79")
    r += 1
    ws.cell(row=r, column=2,
            value="Grundlage der MD-Stichprobenziehung am Prüfungstag."
            ).font = Font(italic=True, color="595959")
    r += 2

    def head(text):
        nonlocal r
        c = ws.cell(row=r, column=2, value=text)
        c.font = Font(bold=True, size=12, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3F3F46")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1

    def line(text, indent=0, bold=False):
        nonlocal r
        c = ws.cell(row=r, column=2, value=("    " * indent) + text)
        c.font = Font(bold=bold, size=10)
        c.alignment = Alignment(wrap_text=False, vertical="top")
        r += 1

    head("So wird die Liste geführt")
    r += 1
    for txt in [
        "1. Pro Klient eine Zeile im Blatt „Klientenliste“.",
        "2. Die gelb markierte Beispielzeile (Zeile 3) vor der ersten Nutzung "
        "löschen.",
        "3. Grau hinterlegte Spalten („Kategorie“, „Herleitung“) sind Formeln "
        "– nicht überschreiben.",
        "4. Felder mit Dropdown nur über die Auswahlliste befüllen, sonst "
        "greift die Kategorie-Logik nicht.",
        "5. Spalte „Stand“ bei jeder Änderung aktualisieren – der MD fragt "
        "nach der Aktualität der Liste.",
        "6. Rote Ecke an einer Überschrift = Ausfüllhinweis, mit der Maus "
        "darüberfahren.",
    ]:
        line(txt)
    r += 1

    head("Kategorie A–D: Matrix aus Mobilität und Kognition")
    r += 1
    line("Die Kategorie wird automatisch aus den beiden Spalten „Mobilität“ "
         "und „Kognition / Kommunikation“ berechnet:", bold=True)
    r += 1

    matrix_top = r
    hdr = ["", "Kognition: uneingeschränkt", "Kognition: eingeschränkt",
           "Kognition: komplett eingeschr."]
    for i, h in enumerate(hdr):
        c = ws.cell(row=r, column=2 + i, value=h)
        c.font = Font(bold=True, size=10,
                      color="FFFFFF" if i else "1A1A1A")
        c.fill = PatternFill("solid", fgColor="5B3A78" if i else "FFFFFF")
        c.alignment = Alignment(wrap_text=True, horizontal="center",
                                vertical="center")
        c.border = BORDER
    ws.row_dimensions[r].height = 32
    r += 1

    rows = [("Mobilität: uneingeschränkt", ["A", "B", "C"]),
            ("Mobilität: eingeschränkt", ["B", "B", "C"]),
            ("Mobilität: komplett eingeschränkt", ["C", "C", "D"])]
    colors = {"A": ("C6EFCE", "1E6B33"), "B": ("FFEB9C", "8A6100"),
              "C": ("FFCC99", "9C4A16"), "D": ("FFC7CE", "9C0006")}
    for label, vals in rows:
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B3A78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        for i, v in enumerate(vals):
            bg, fg = colors[v]
            cc = ws.cell(row=r, column=3 + i, value=v)
            cc.font = Font(bold=True, size=13, color=fg)
            cc.fill = PatternFill("solid", fgColor=bg)
            cc.alignment = Alignment(horizontal="center", vertical="center")
            cc.border = BORDER
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1

    for txt in [
        "A  =  Mobilität uneingeschränkt UND Kognition uneingeschränkt",
        "B  =  Teileinschränkung bei Mobilität oder Kognition",
        "C  =  Mobilität ODER Kognition komplett eingeschränkt",
        "D  =  Mobilität UND Kognition komplett eingeschränkt "
        "(höchste Komplexitätsstufe)",
    ]:
        line(txt)
    r += 1
    c = ws.cell(row=r, column=2,
                value="Hinweis zur Auslegung: Die Kombinationen „komplett "
                      "eingeschränkt“ + „eingeschränkt“ erfüllen wörtlich "
                      "sowohl Regel B als auch Regel C.")
    c.font = Font(size=10, color="9C4A16")
    r += 1
    c = ws.cell(row=r, column=2,
                value="Die Vorlage wertet in diesen Fällen die schwerere "
                      "Einstufung (C), da die Kategorie den höchsten "
                      "Komplexitätsgrad abbilden soll.")
    c.font = Font(size=10, color="9C4A16")
    r += 2

    head("Verteilung (aktualisiert sich automatisch)")
    r += 1
    stats_row = r
    for i, (label, formula) in enumerate([
        ("Kategorie A", 'A'), ("Kategorie B", 'B'),
        ("Kategorie C", 'C'), ("Kategorie D", 'D')]):
        c = ws.cell(row=stats_row + i, column=2, value=label)
        c.font = Font(bold=True, size=10)
        bg, fg = colors[formula]
        c.fill = PatternFill("solid", fgColor=bg)
        c.border = BORDER
        v = ws.cell(row=stats_row + i, column=3,
                    value=f'=COUNTIF(Klientenliste!${col_kat}${DATA_START}:'
                          f'${col_kat}${DATA_END},"{formula}")')
        v.font = Font(bold=True, size=11, color=fg)
        v.alignment = Alignment(horizontal="center")
        v.fill = PatternFill("solid", fgColor=bg)
        v.border = BORDER
    r = stats_row + 4
    c = ws.cell(row=r, column=2, value="Klienten gesamt")
    c.font = Font(bold=True, size=10)
    c.border = BORDER
    v = ws.cell(row=r, column=3,
                value=f'=COUNTA(Klientenliste!$A${DATA_START}:$A${DATA_END})')
    v.font = Font(bold=True, size=11)
    v.alignment = Alignment(horizontal="center")
    v.border = BORDER
    r += 2

    head("Datenschutz")
    r += 1
    for txt in [
        "Diese Liste enthält Gesundheitsdaten – besondere Kategorien "
        "personenbezogener Daten nach Art. 9 DSGVO.",
        "• Nur auf dienstlichen, zugriffsgeschützten Systemen speichern.",
        "• Nicht unverschlüsselt per E-Mail oder Messenger versenden.",
        "• Ausdrucke nach der Prüfung sicher vernichten.",
        "• Zugriff auf den Kreis beschränken, der die Liste pflegen muss.",
    ]:
        line(txt)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "docs/internal/anforderungen/QPR-Klientenliste-2026.xlsx"
    if len(sys.argv) > 2:
        DATA_ROWS = int(sys.argv[2])
        DATA_END = DATA_START + DATA_ROWS - 1
    print("geschrieben:", build(out), f"({DATA_ROWS} Zeilen)")
